#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Build a visual-heavy customer proposal deck for GloSource AI."""

from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path("docs/business/out")
ASSETS = OUT / "visual-proposal-assets"
IMG = Path("public/images")
PPTX = OUT / "05-사업제안서_visual.pptx"

W, H = 1600, 900
NAVY = "0D1B2A"
INK = RGBColor(0x0D, 0x1B, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
MINERAL = RGBColor(0x8F, 0x3A, 0x26)
GREEN = RGBColor(0x0F, 0x7A, 0x55)
RED = RGBColor(0xB4, 0x23, 0x18)
AMBER = RGBColor(0xB7, 0x6E, 0x00)
PAPER = RGBColor(0xF8, 0xF5, 0xEF)


def font(size: int, bold: bool = False):
    names = [
        "C:/Windows/Fonts/malgunbd.ttf" if bold else "C:/Windows/Fonts/malgun.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for name in names:
        if Path(name).exists():
            return ImageFont.truetype(name, size)
    return ImageFont.load_default()


def save_diagram(name: str, title: str, items: list[tuple[str, str, str]], footer: str = "") -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    path = ASSETS / name
    im = Image.new("RGB", (W, H), "#F8F5EF")
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((48, 48, W - 48, H - 48), radius=36, fill="#FFFFFF", outline="#E2DDD2", width=3)
    d.text((90, 86), title, fill=f"#{NAVY}", font=font(54, True))
    x = 105
    card_w = 420
    for i, (head, body, color) in enumerate(items):
        left = x + i * 490
        d.rounded_rectangle((left, 230, left + card_w, 635), radius=28, fill=color, outline=color)
        d.ellipse((left + 28, 258, left + 108, 338), fill="#FFFFFF")
        d.text((left + 56, 271), str(i + 1), fill=color, font=font(44, True), anchor="ma")
        d.text((left + 40, 375), head, fill="#FFFFFF", font=font(38, True))
        wrapped = wrap(body, 18)
        d.multiline_text((left + 40, 450), wrapped, fill="#FFFFFF", font=font(25), spacing=10)
        if i < len(items) - 1:
            d.line((left + card_w + 30, 432, left + card_w + 78, 432), fill="#C7BFAF", width=8)
            d.polygon([(left + card_w + 78, 432), (left + card_w + 52, 414), (left + card_w + 52, 450)], fill="#C7BFAF")
    if footer:
        d.text((90, 735), footer, fill="#64748B", font=font(30, True))
    im.save(path)
    return path


def wrap(text: str, max_chars: int) -> str:
    words = text.split()
    lines, cur = [], ""
    for word in words:
        if len(cur) + len(word) + 1 > max_chars:
            lines.append(cur)
            cur = word
        else:
            cur = f"{cur} {word}".strip()
    if cur:
        lines.append(cur)
    return "\n".join(lines)


def make_assets():
    save_diagram(
        "three-step.png",
        "AI가 찾고, 정부 원장으로 확인합니다",
        [
            ("AI 발굴", "웹에서 한국 공급사 후보를 수집", "#264653"),
            ("원장 대조", "식약처·HACCP·국세청 데이터로 확인", "#8F3A26"),
            ("판정 표시", "검증·위험·미검증을 분리", "#0F7A55"),
        ],
        "GloSource는 소싱 대행이 아니라 검증 리서치 도구입니다.",
    )
    save_diagram(
        "risk-cards.png",
        "바이어가 가장 많이 놓치는 3가지",
        [
            ("중개상", "공장처럼 보이지만 남의 공장을 재판매", "#B42318"),
            ("미등록", "제조업 원장에 없는 업체", "#B76E00"),
            ("폐업", "웹사이트는 살아있지만 사업자는 중단", "#5A4A42"),
        ],
        "잘못 접촉하면 샘플비·리드타임·품질책임이 모두 흔들립니다.",
    )
    save_diagram(
        "registry-stack.png",
        "검증은 광고 문구가 아니라 원장 데이터로 합니다",
        [
            ("MFDS", "화장품 제조/책임판매 구분", "#244C5A"),
            ("HACCP", "식품 안전 인증 원장 대조", "#7B3F25"),
            ("NTS", "계속·휴업·폐업 실시간 확인", "#0F6650"),
        ],
        "검증할 수 없는 카테고리는 검증됨으로 표시하지 않습니다.",
    )
    save_diagram(
        "pilot-flow.png",
        "무료 파일럿은 이렇게 진행됩니다",
        [
            ("제품 브리프", "카테고리·MOQ·인증 요건 전달", "#264653"),
            ("검증 리포트", "후보·원장대조·위험신호 정리", "#8F3A26"),
            ("피드백", "쓸 만한지 솔직하게 확인", "#0F7A55"),
        ],
        "비용 0원. 신용카드 없음. 피드백 1회만 요청합니다.",
    )


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_bg(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_textbox(slide, text, x, y, w, h, size=24, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    r = p.runs[0]
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def title(slide, text, subtitle=None):
    add_textbox(slide, text, 0.65, 0.42, 8.2, 0.65, 30, True, INK)
    if subtitle:
        add_textbox(slide, subtitle, 0.68, 1.08, 8.2, 0.35, 11, False, MUTED)


def add_image(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_chip(slide, text, x, y, w, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.46))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.color.rgb = color
    add_textbox(slide, text, x + 0.08, y + 0.03, w - 0.16, 0.36, 12, True, RGBColor(255, 255, 255), PP_ALIGN.CENTER)


def add_metric(slide, value, label, x, y, color=MINERAL):
    add_textbox(slide, value, x, y, 2.1, 0.56, 26, True, color, PP_ALIGN.CENTER)
    add_textbox(slide, label, x, y + 0.56, 2.1, 0.38, 10, False, MUTED, PP_ALIGN.CENTER)


def build():
    make_assets()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 Cover
    s = blank_slide(prs)
    add_bg(s)
    add_image(s, IMG / "glosource-hero-command-room.png", 5.55, 0, 7.78, 7.5)
    add_image(s, IMG / "glosource-logo-transparent.png", 0.75, 0.62, 1.25, 1.25)
    add_textbox(s, "GloSource AI", 0.75, 2.1, 4.6, 0.7, 36, True)
    add_textbox(s, "한국 공급사를 처음 찾는 해외 바이어에게\n“진짜 공장인지” 먼저 확인해주는 리서치 도구", 0.78, 2.95, 4.7, 1.1, 21, True, MINERAL)
    add_textbox(s, "Business Proposal · glosourceai.com · hello@glosourceai.com", 0.8, 6.75, 4.8, 0.35, 11, False, MUTED)

    # 2 Pain
    s = blank_slide(prs)
    add_bg(s)
    title(s, "검색 상위가 곧 공장은 아닙니다", "Buyer problem")
    add_image(s, ASSETS / "risk-cards.png", 0.65, 1.35, 7.4, 4.15)
    add_image(s, IMG / "k-product-grid.png", 8.45, 1.05, 4.25, 4.85)
    add_textbox(s, "바이어는 샘플비를 보내기 전\n제조사·인증·폐업 여부를 확인해야 합니다.", 8.55, 6.05, 4.05, 0.85, 18, True, INK)

    # 3 Market
    s = blank_slide(prs)
    add_bg(s)
    title(s, "문제는 커지고 있습니다", "K-Beauty export demand is expanding, and small buyers are entering Korea for the first time.")
    add_image(s, IMG / "k-wave-demand.png", 0.7, 1.35, 5.5, 4.8)
    add_metric(s, "$11.4B", "2025 한국 화장품 수출", 6.65, 1.75)
    add_metric(s, "#2", "세계 화장품 수출 순위", 9.2, 1.75)
    add_metric(s, "69.8%", "중소기업 수출 비중", 6.65, 3.55)
    add_metric(s, "90초", "검증 shortlist 목표", 9.2, 3.55)
    add_textbox(s, "대형 유통사는 이미 공급망이 있습니다.\nGloSource는 처음 한국을 찾는 소규모 바이어를 위한 도구입니다.", 6.75, 5.45, 5.25, 0.95, 19, True, INK)

    # 4 Solution
    s = blank_slide(prs)
    add_bg(s)
    title(s, "AI 추천이 아니라, AI + 정부 원장 대조입니다", "Discovery and verification are separated.")
    add_image(s, ASSETS / "three-step.png", 0.7, 1.35, 11.9, 5.25)

    # 5 Workflow
    s = blank_slide(prs)
    add_bg(s)
    title(s, "고객은 제품 조건만 입력합니다", "Then GloSource turns it into a verified shortlist.")
    add_image(s, IMG / "workflow-product-input.png", 0.65, 1.35, 3.95, 4.2)
    add_image(s, IMG / "workflow-supplier-analysis.png", 4.7, 1.35, 3.95, 4.2)
    add_image(s, IMG / "workflow-rfq-match.png", 8.75, 1.35, 3.95, 4.2)
    add_chip(s, "1. Brief", 1.55, 5.9, 2.1, MINERAL)
    add_chip(s, "2. Verify", 5.6, 5.9, 2.1, GREEN)
    add_chip(s, "3. RFQ", 9.65, 5.9, 2.1, MINERAL)

    # 6 Registry stack
    s = blank_slide(prs)
    add_bg(s)
    title(s, "검증의 기준은 정부 원장입니다", "No badge without registry evidence.")
    add_image(s, ASSETS / "registry-stack.png", 0.72, 1.35, 7.3, 4.7)
    add_image(s, IMG / "glosource-sourcing-workspace.png", 8.35, 1.1, 4.45, 4.9)
    add_textbox(s, "검증 가능 / 위험 / 미검증을 분리 표시", 8.45, 6.1, 4.15, 0.52, 20, True, INK)

    # 7 What buyer gets
    s = blank_slide(prs)
    add_bg(s)
    title(s, "결과물은 바로 의사결정에 쓰는 리포트입니다", "Shortlist, evidence, red flags, RFQ draft.")
    add_image(s, IMG / "glosource-enterprise-ops.png", 0.7, 1.25, 6.1, 4.95)
    for i, text in enumerate(["제조사 vs 중개상", "CGMP / HACCP 대조", "국세청 영업상태", "MOQ·리드타임 비교", "영문 RFQ 초안"]):
        add_chip(s, text, 7.25, 1.45 + i * 0.82, 4.4, [GREEN, MINERAL, RED, AMBER, INK][i])
    add_textbox(s, "고객이 직접 한국어 원장을 뒤지는 시간을 줄이고,\n위험한 후보를 먼저 걸러냅니다.", 7.3, 6.0, 4.8, 0.75, 18, True, INK)

    # 8 Honest boundaries
    s = blank_slide(prs)
    add_bg(s)
    title(s, "하지 않는 것을 분명히 합니다", "Trust comes from boundaries.")
    add_image(s, IMG / "k-supplier-process.png", 0.75, 1.25, 5.55, 4.85)
    rows = [
        ("소싱 대행", "하지 않음"),
        ("공장 감사", "하지 않음"),
        ("물류·거래 중개", "하지 않음"),
        ("검증 불가를 검증됨 표시", "하지 않음"),
    ]
    for i, (a, b) in enumerate(rows):
        y = 1.6 + i * 0.95
        add_textbox(s, a, 7.0, y, 3.2, 0.35, 18, True, INK)
        add_chip(s, b, 10.35, y - 0.05, 1.45, RED)
    add_textbox(s, "우리는 공급사를 팔지 않습니다.\n확인 가능한 사실을 팝니다.", 7.0, 5.8, 4.8, 0.75, 22, True, MINERAL)

    # 9 Free pilot
    s = blank_slide(prs)
    add_bg(s)
    title(s, "처음 고객에게는 무료 파일럿부터 제안합니다", "No card. No fake promise. One honest feedback loop.")
    add_image(s, ASSETS / "pilot-flow.png", 0.7, 1.35, 8.0, 4.65)
    add_image(s, IMG / "k-hero-products.png", 9.1, 1.25, 3.55, 3.9)
    add_textbox(s, "무료 리포트 1건\n→ 솔직한 피드백 1회", 9.25, 5.45, 3.1, 0.9, 22, True, INK, PP_ALIGN.CENTER)

    # 10 Pricing
    s = blank_slide(prs)
    add_bg(s)
    title(s, "유료 전환은 검증 후에", "Simple subscription, already wired through Polar.")
    prices = [("Free", "$0", "10 searches"), ("Starter", "$39", "40 searches"), ("Pro", "$99", "100 searches"), ("Business", "$249", "300 searches")]
    for i, (plan, price, quota) in enumerate(prices):
        x = 0.85 + i * 3.1
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(2.55), Inches(3.55))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shp.line.color.rgb = RGBColor(0xE2, 0xDD, 0xD2)
        add_textbox(s, plan, x + 0.25, 2.15, 2.05, 0.35, 18, True, INK, PP_ALIGN.CENTER)
        add_textbox(s, price, x + 0.25, 2.85, 2.05, 0.62, 31, True, MINERAL, PP_ALIGN.CENTER)
        add_textbox(s, quota, x + 0.25, 3.72, 2.05, 0.35, 15, False, MUTED, PP_ALIGN.CENTER)
        add_textbox(s, "per month", x + 0.25, 4.32, 2.05, 0.25, 10, False, MUTED, PP_ALIGN.CENTER)
    add_textbox(s, "현재 고객 레퍼런스는 없습니다. 그래서 먼저 무료 파일럿으로 검증합니다.", 1.15, 6.15, 10.9, 0.45, 18, True, INK, PP_ALIGN.CENTER)

    # 11 Partner proposal
    s = blank_slide(prs)
    add_bg(s)
    title(s, "기관·플랫폼에는 검증 레이어를 제공합니다", "For trade agencies, accelerators, B2B platforms.")
    add_image(s, IMG / "glosource-hero-command-room.png", 0.7, 1.25, 5.3, 4.85)
    partner_items = [("공동 마케팅", "한국 공급사 검증 콘텐츠"), ("회원사 제공", "기관 단위 라이선스"), ("데이터 협력", "원장 대조 결과 피드백"), ("플랫폼 임베드", "API/위젯 연동")]
    for i, (head, body) in enumerate(partner_items):
        x = 6.55 + (i % 2) * 2.95
        y = 1.55 + (i // 2) * 2.05
        add_chip(s, head, x, y, 2.35, GREEN if i % 2 == 0 else MINERAL)
        add_textbox(s, body, x, y + 0.62, 2.55, 0.55, 16, True, INK)
    add_textbox(s, "첫 단계는 대규모 계약이 아니라 소수 회원사 무료 파일럿입니다.", 6.65, 5.95, 5.2, 0.5, 18, True, INK)

    # 12 CTA
    s = blank_slide(prs)
    add_bg(s)
    add_image(s, IMG / "glosource-logo-transparent.png", 5.65, 0.72, 2.0, 2.0)
    add_textbox(s, "무료 공급사 검증 파일럿을 시작하세요", 1.35, 3.0, 10.7, 0.65, 34, True, INK, PP_ALIGN.CENTER)
    add_textbox(s, "제품 카테고리 · MOQ · 인증 요건을 보내주시면\n검증 shortlist와 RFQ 초안을 먼저 드립니다.", 2.35, 3.9, 8.7, 0.9, 22, True, MINERAL, PP_ALIGN.CENTER)
    add_textbox(s, "hello@glosourceai.com", 3.75, 5.45, 5.8, 0.55, 28, True, GREEN, PP_ALIGN.CENTER)
    add_textbox(s, "glosourceai.com", 4.55, 6.1, 4.2, 0.35, 18, False, MUTED, PP_ALIGN.CENTER)

    prs.save(PPTX)
    print(f"saved {PPTX}")


if __name__ == "__main__":
    build()
